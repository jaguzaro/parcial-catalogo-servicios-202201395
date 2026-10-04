#!/usr/bin/env python3
"""Analiza data/CatalogoServicios.xlsx y escribe docs/contexto/analisis-excel.md.

Solo lectura sobre el Excel. El informe es descriptivo y determinista: no incluye fecha
ni hora, y todas las listas salen ordenadas, así que dos ejecuciones dan el mismo archivo.

Uso (desde la raíz del proyecto, dentro del contenedor):
    python scripts/analisis/analizar_excel.py
"""
import hashlib
import re
import sys
from collections import OrderedDict, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries

RAIZ = Path(__file__).resolve().parents[2]
XLSX = RAIZ / "data" / "CatalogoServicios.xlsx"
ENUNCIADO = RAIZ / "docs" / "contexto" / "enunciado.md"
SALIDA = RAIZ / "docs" / "contexto" / "analisis-excel.md"

COMANDO = (
    'docker run --rm -v "$PWD":/work -v "$PWD/data":/work/data:ro -w /work python:3.12-slim \\\n'
    '  sh -c "pip install --quiet openpyxl && python scripts/analisis/analizar_excel.py"'
)

HOJA = "Servicios Externos"
FILA_ENCABEZADO = 4
PRIMERA, ULTIMA = 5, 101          # filas de datos según el enunciado (sección 2)
OPCIONES = {"E": "ACTIVO", "F": "CLASE DE SERVICIO", "G": "CRITICIDAD", "H": "TIPO DE SERVICIO"}
OPC_FILAS = (112, 122)
ESPERADO_N1, ESPERADO_N2 = 12, 46

# Patrones para la sección final. "revel" sale de una revisión manual previa de I5.
PATRONES_SOSPECHOSOS = [
    r"ignor", r"ejecut", r"ahora eres", r"olvida", r"instrucci", r"prompt", r"system",
    r"claude", r"asistente", r"\bia\b", r"\bllm\b", r"revel", r"responde", r"borra",
    r"elimina", r"commit", r"push", r"https?://",
]


# ---------------------------------------------------------------- utilidades
def rep(v):
    """Valor literal entre backticks (repr conserva espacios finales y tipos)."""
    if v is None:
        return "*(vacía)*"
    return "`" + repr(v).replace("|", "\\|").replace("\n", "\\n") + "`"


def tabla(encabezados, filas):
    out = ["| " + " | ".join(encabezados) + " |", "|" + "|".join("---" for _ in encabezados) + "|"]
    for f in filas:
        out.append("| " + " | ".join(str(x).replace("\n", " ") for x in f) + " |")
    return "\n".join(out)


def rangos_compactos(nums):
    """[5,6,7,9] -> '5–7, 9'."""
    nums = sorted(nums)
    partes, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        partes.append(str(nums[i]) if i == j else f"{nums[i]}–{nums[j]}")
        i = j + 1
    return ", ".join(partes)


def vacio(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def tipo(v):
    if v is None:
        return "vacío"
    if isinstance(v, bool):
        return "booleano"
    if isinstance(v, (int, float)):
        return "numérico"
    if isinstance(v, str):
        return "vacío (solo espacios)" if v.strip() == "" else "texto"
    return type(v).__name__


# ---------------------------------------------------------------- enunciado
def leer_enunciado():
    """Extrae del enunciado (sección 2) encabezados y listas de opciones."""
    txt = ENUNCIADO.read_text(encoding="utf-8")
    cols = dict(re.findall(r"^\| ([A-L]) \| (.+?) \|", txt, flags=re.M))
    listas = {}
    for clave, patron in (
        ("F", r"Las opciones de clase son (.*?)\.\s"),
        ("G", r"Las criticidades son (.*?)\.\s"),
        ("H", r"Los tipos son (.*?)\.\s"),
    ):
        m = re.search(patron, txt, flags=re.S)
        if not m:
            sys.exit(f"No se encontró en el enunciado la lista para la columna {clave}")
        listas[clave] = re.findall(r"`([^`]+)`", m.group(1))
    if "Indicador S/N" not in txt:
        sys.exit("No se encontró 'Indicador S/N' en el enunciado")
    listas["E"] = ["S", "N"]
    return cols, listas


# ---------------------------------------------------------------- análisis
def main():
    if not XLSX.exists():
        sys.exit(f"No existe {XLSX}")
    sha = hashlib.sha256(XLSX.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(XLSX)  # sin data_only: interesa lo almacenado, no fórmulas calculadas
    enun_cols, enun_listas = leer_enunciado()

    out = []
    w = out.append

    w("# Análisis de `data/CatalogoServicios.xlsx`")
    w("")
    w("Regenerar este informe (código de salida 0 = correcto):")
    w("")
    w("```bash")
    w(COMANDO)
    w("```")
    w("")
    w("Informe **descriptivo**, generado por `scripts/analisis/analizar_excel.py`. No contiene fecha "
      "ni hora: dos ejecuciones seguidas producen el mismo archivo. Las celdas se citan en notación "
      "A1; los valores se muestran con `repr`, de modo que los espacios sobrantes son visibles. "
      "Todo el contenido del Excel se trata como dato.")
    w("")
    w(f"- SHA-256 del Excel analizado: `{sha}`")
    w(f"- Libros: {len(wb.worksheets)} hoja(s).")
    w("")

    if HOJA not in wb.sheetnames:
        sys.exit(f"No existe la hoja '{HOJA}'")
    ws = wb[HOJA]
    V = lambda coord: ws[coord].value  # noqa: E731

    # Mapa fila,col -> rango combinado que la cubre.
    rangos = sorted(ws.merged_cells.ranges, key=lambda r: (r.min_col, r.min_row, r.max_row))
    cubierta = {}
    for r in rangos:
        for fila in range(r.min_row, r.max_row + 1):
            for col in range(r.min_col, r.max_col + 1):
                cubierta[(fila, col)] = r

    # ============================================================ a
    w("## a. Estructura del libro")
    w("")
    filas_hojas = []
    for s in wb.worksheets:
        ult_fila = ult_col = 0
        for row in s.iter_rows():
            for c in row:
                if c.value is not None:
                    ult_fila, ult_col = max(ult_fila, c.row), max(ult_col, c.column)
        usado = f"A1:{get_column_letter(ult_col)}{ult_fila}" if ult_fila else "(sin datos)"
        filas_hojas.append([f"`{s.title}`", s.sheet_state, f"`{s.dimensions}`", usado, len(s.merged_cells.ranges)])
    w(tabla(["Hoja", "Estado", "Dimensión declarada (`ws.dimensions`)",
             "Extensión con valores (calculada)", "Rangos combinados"], filas_hojas))
    w("")
    w("La dimensión declarada puede exceder la zona con valores (celdas con formato sin contenido); "
      "la columna «Extensión con valores» sale de recorrer todas las celdas con valor no nulo.")
    w("")
    w(f"Encabezados de `A{FILA_ENCABEZADO}:L{FILA_ENCABEZADO}` (texto literal) y comparación con la tabla "
      "de columnas del enunciado (sección 2):")
    w("")
    filas = []
    for i in range(1, 13):
        letra = get_column_letter(i)
        lit = V(f"{letra}{FILA_ENCABEZADO}")
        esp = enun_cols.get(letra)
        filas.append([f"`{letra}{FILA_ENCABEZADO}`", rep(lit), rep(esp), "coincide" if lit == esp else "**difiere**"])
    w(tabla(["Celda", "Texto en el Excel", "Campo según el enunciado", "Resultado"], filas))
    w("")
    resto = [c.coordinate for c in ws[FILA_ENCABEZADO] if c.column > 12 and c.value is not None]
    w("Celdas con valor en la fila 4 fuera de `A:L`: " + (", ".join(resto) if resto else "ninguna") + ".")
    fuera = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            r_, k_ = c.row, c.column
            en_datos = FILA_ENCABEZADO <= r_ <= ULTIMA and k_ <= 12
            en_opc = (OPC_FILAS[0] - 1) <= r_ <= OPC_FILAS[1] and 5 <= k_ <= 8
            if not (en_datos or en_opc):
                fuera.append(c)
    w("")
    w("Celdas con valor fuera de la tabla (`A4:L101`) y del bloque de opciones (`E111:H122`):")
    w("")
    w(tabla(["Celda", "Valor"], [[f"`{c.coordinate}`", rep(c.value)] for c in fuera]) if fuera else "Ninguna.")
    w("")

    # ============================================================ b
    w("## b. Celdas combinadas")
    w("")
    w(f"Total: **{len(rangos)}** rangos en `{HOJA}`. Valor de la celda principal (esquina superior "
      "izquierda) de cada uno; las demás celdas del rango se leen vacías en openpyxl.")
    w("")
    filas = []
    for r in rangos:
        ancla = ws.cell(row=r.min_row, column=r.min_col)
        filas.append([f"`{r.coord}`", f"{r.max_row - r.min_row + 1}", f"`{ancla.coordinate}`", rep(ancla.value)])
    w(tabla(["Rango", "Filas", "Celda principal", "Valor de la celda principal"], filas))
    w("")
    por_col = defaultdict(int)
    for r in rangos:
        por_col[get_column_letter(r.min_col)] += 1
    w("Rangos por columna: " + ", ".join(f"{k}: {por_col[k]}" for k in sorted(por_col)) + ".")
    multi = [r.coord for r in rangos if r.min_col != r.max_col]
    w("Rangos que abarcan más de una columna: " + (", ".join(multi) if multi else "ninguno") + ".")
    fuera_zona = [r.coord for r in rangos if r.min_row < PRIMERA or r.max_row > ULTIMA]
    w(f"Rangos fuera de las filas {PRIMERA}–{ULTIMA}: " + (", ".join(fuera_zona) if fuera_zona else "ninguno") + ".")
    w("")

    # ============================================================ c
    w("## c. Conteos de control")
    w("")
    n1 = OrderedDict()   # código -> [celdas]
    n2 = OrderedDict()
    for f in range(PRIMERA, ULTIMA + 1):
        a, c_ = V(f"A{f}"), V(f"C{f}")
        if not vacio(a):
            n1.setdefault(a.strip() if isinstance(a, str) else a, []).append(f"A{f}")
        if not vacio(c_):
            n2.setdefault(c_.strip() if isinstance(c_, str) else c_, []).append(f"C{f}")
    cn1, cn2 = len(n1), len(n2)
    w(tabla(["Control", "Observado", "Esperado (enunciado, sección 2)", "Resultado"], [
        [f"Códigos de nivel 1 distintos (columna A, filas {PRIMERA}–{ULTIMA})", cn1, ESPERADO_N1,
         "**coincide**" if cn1 == ESPERADO_N1 else "**difiere**"],
        [f"Códigos de nivel 2 explícitos distintos (columna C, filas {PRIMERA}–{ULTIMA})", cn2, ESPERADO_N2,
         "**coincide**" if cn2 == ESPERADO_N2 else "**difiere**"],
    ]))
    w("")
    w("«Explícito» = código escrito en una celda con valor (en un rango combinado, su celda principal). "
      "Criterio aplicado tal cual, sin ajustes para que cuadre.")
    w("")
    w("Celdas de la columna A con valor, y total de celdas frente a códigos distintos: "
      f"{sum(len(v) for v in n1.values())} celdas / {cn1} códigos.")
    w("")
    w(tabla(["Código N1", "Celdas (columna A)", "Nº celdas"],
            [[f"`{k}`", ", ".join(f"`{c}`" for c in v), len(v)] for k, v in n1.items()]))
    w("")
    w(f"Celdas de la columna C con valor: {sum(len(v) for v in n2.values())} celdas / {cn2} códigos.")
    w("")
    w(tabla(["Código N2", "Celdas (columna C)", "Nº celdas"],
            [[f"`{k}`", ", ".join(f"`{c}`" for c in v), len(v)] for k, v in n2.items()]))
    w("")
    rep_n2 = [k for k, v in n2.items() if len(v) > 1]
    w("Códigos de nivel 2 repetidos: " + (", ".join(rep_n2) if rep_n2 else "ninguno") + ".")
    patron_n2 = re.compile(r"^SE\.\d{2}\.\d{2}$")
    raros = [(k, v) for k, v in n2.items() if not patron_n2.match(str(k))]
    w("Códigos de nivel 2 que no siguen el patrón `SE.NN.NN` (dos dígitos en cada tramo): " +
      (", ".join(f"`{k}` ({v[0]})" for k, v in raros) if raros else "ninguno") + ".")
    tipos_cod = sorted({tipo(ws[c].value) for v in list(n1.values()) + list(n2.values()) for c in v})
    w("Tipos de dato de esas celdas de código: " + ", ".join(tipos_cod) + ".")
    con_espacios = [c for v in list(n1.values()) + list(n2.values()) for c in v
                    if isinstance(ws[c].value, str) and ws[c].value != ws[c].value.strip()]
    w("Códigos con espacios al inicio o al final: " + (", ".join(con_espacios) if con_espacios else "ninguno") + ".")
    w("")
    # coherencia prefijo N2 con el N1 vigente (ancla A más cercana hacia arriba)
    w("Coherencia del prefijo de cada código N2 con el código N1 vigente en su fila "
      "(valor de la celda A principal más cercana hacia arriba, dentro de las filas de datos):")
    incoh = []
    vigente = None
    for f in range(PRIMERA, ULTIMA + 1):
        if not vacio(V(f"A{f}")):
            vigente = (V(f"A{f}").strip(), f"A{f}")
        c_ = V(f"C{f}")
        if not vacio(c_) and vigente and not str(c_).strip().startswith(vigente[0] + "."):
            incoh.append(f"C{f} `{c_}` vs {vigente[1]} `{vigente[0]}`")
    w("")
    w("Incoherencias: " + ("; ".join(incoh) if incoh else "ninguna") + ".")
    w("")

    # ============================================================ d
    w("## d. Filas 99 a 101")
    w("")
    w("Estado de cada celda de `A:L`. Una celda cuenta como vacía si es `None` o solo tiene espacios.")
    w("")
    nombres = {get_column_letter(i): V(f"{get_column_letter(i)}{FILA_ENCABEZADO}") for i in range(1, 13)}
    for f in (99, 100, 101):
        w(f"### Fila {f}")
        w("")
        filas, vacias = [], []
        for i in range(1, 13):
            L = get_column_letter(i)
            v = V(f"{L}{f}")
            est = "vacía" if vacio(v) else "con valor"
            if vacio(v):
                vacias.append(f"{L}{f}")
            filas.append([f"`{L}{f}`", nombres[L], est, rep(v) if not vacio(v) else ""])
        w(tabla(["Celda", "Columna", "Estado", "Valor"], filas))
        w("")
        w(f"Vacías ({len(vacias)}): " + ", ".join(vacias) + ".")
        w("")
    w("Resumen por columna (filas 99, 100, 101):")
    w("")
    filas = []
    for i in range(1, 13):
        L = get_column_letter(i)
        filas.append([f"{L} — {nombres[L]}"] + ["vacía" if vacio(V(f"{L}{f}")) else "con valor" for f in (99, 100, 101)])
    w(tabla(["Columna", "Fila 99", "Fila 100", "Fila 101"], filas))
    w("")
    ref = [f for f in range(PRIMERA, 99)]
    con_efgh = [f for f in ref if all(not vacio(V(f"{L}{f}")) for L in "EFGH")]
    w(f"Para contraste: de las filas {PRIMERA}–98, {len(con_efgh)} de {len(ref)} tienen las cuatro columnas "
      "E, F, G y H con valor.")
    w("")

    # ============================================================ e
    w("## e. Código `SE.12`")
    w("")
    hits = []
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "SE.12" in c.value:
                hits.append(c)
    w("Todas las celdas de la hoja cuyo texto contiene `SE.12`:")
    w("")
    w(tabla(["Celda", "Valor", "Coincidencia exacta con `SE.12`"],
            [[f"`{c.coordinate}`", rep(c.value), "sí" if c.value.strip() == "SE.12" else "no (prefijo de un código N2)"]
             for c in hits]))
    w("")
    exactas = [c for c in hits if c.value.strip() == "SE.12"]
    filas_se12 = sorted({c.row for c in hits})
    w(f"Celdas con `SE.12` exacto: {', '.join(c.coordinate for c in exactas)} ({len(exactas)}). "
      f"Filas donde aparece algún código `SE.12*`: {rangos_compactos(filas_se12)}.")
    w("")
    w("Valor de cada columna en esas filas:")
    w("")
    filas = []
    for i in range(1, 13):
        L = get_column_letter(i)
        filas.append([f"{L} — {nombres[L]}"] + [rep(V(f"{L}{f}")) for f in filas_se12])
    w(tabla(["Columna"] + [f"Fila {f}" for f in filas_se12], filas))
    w("")
    nombres_n1 = sorted({c.offset(0, 1).value for c in exactas}, key=str)
    w("Valores de `B` en las filas con `SE.12` exacto: " + "; ".join(
        f"`B{c.row}` = {rep(c.offset(0, 1).value)}" for c in exactas) + ".")
    w("Los nombres " + (f"**difieren** ({len(nombres_n1)} nombres distintos)" if len(nombres_n1) > 1
                         else "coinciden") + " para el mismo código N1.")
    w("")
    w("Nombres de nivel 2 en esas filas (columna D): " + "; ".join(
        f"`D{f}` = {rep(V(f'D{f}'))}" for f in filas_se12) + ".")
    w("")
    # nombres N1 repetidos en otros códigos y nombres N1/N2 iguales
    n1_nombres = defaultdict(set)
    for f in range(PRIMERA, ULTIMA + 1):
        if not vacio(V(f"A{f}")):
            n1_nombres[V(f"A{f}").strip()].add(V(f"B{f}"))
    conf = {k: v for k, v in n1_nombres.items() if len(v) > 1}
    w("Códigos N1 con más de un nombre en la columna B (todos los códigos, no solo `SE.12`): " +
      (", ".join(f"`{k}`" for k in sorted(conf)) if conf else "ninguno") + ".")
    n1_set = {n for v in n1_nombres.values() for n in v}
    coinc = []
    for k, v in n2.items():
        d = ws[f"D{v[0][1:]}"].value
        if d in n1_set:
            coinc.append(f"`D{v[0][1:]}` {rep(d)} (`{k}`)")
    w("Nombres de nivel 2 (columna D) idénticos a algún nombre de nivel 1 (columna B): " +
      ("; ".join(coinc) if coinc else "ninguno") + ".")
    w("")

    # ============================================================ f
    w("## f. Filas sin código fuera de combinaciones")
    w("")
    w(f"Filas {PRIMERA}–{ULTIMA} con la columna C sin valor. Se separan en las que quedan dentro de un rango "
      "combinado de la columna C (continuación del servicio de la celda principal) y las que no.")
    w("")
    sin_codigo = [f for f in range(PRIMERA, ULTIMA + 1) if vacio(V(f"C{f}"))]
    dentro = [f for f in sin_codigo if (f, 3) in cubierta]
    fuera_c = [f for f in sin_codigo if (f, 3) not in cubierta]
    w(f"- Filas con C vacía: {len(sin_codigo)} ({rangos_compactos(sin_codigo)}).")
    w(f"- Dentro de un rango combinado de la columna C: {len(dentro)} ({rangos_compactos(dentro)}).")
    w(f"- **Fuera** de todo rango combinado de la columna C: {len(fuera_c)} ({rangos_compactos(fuera_c) or '—'}).")
    w("")
    if fuera_c:
        w("Contenido de las filas fuera de combinaciones (todas las columnas `A:L`):")
        w("")
        w(tabla(["Fila"] + [f"{get_column_letter(i)}" for i in range(1, 13)],
                [[f] + [rep(V(f"{get_column_letter(i)}{f}")) for i in range(1, 13)] for f in fuera_c]))
        w("")
        w("Contexto observable de cada una (hechos, sin asignarla a ningún servicio):")
        w("")
        for f in fuera_c:
            arriba = next((g for g in range(f - 1, PRIMERA - 1, -1) if not vacio(V(f"C{g}"))), None)
            abajo = next((g for g in range(f + 1, ULTIMA + 1) if not vacio(V(f"C{g}"))), None)
            r_arr = cubierta.get((f - 1, 3))
            r_a, r_b = cubierta.get((f, 1)), cubierta.get((f, 2))
            partes = [f"- Fila {f}: A{f} " + (f"está dentro de `{r_a.coord}` (principal `A{r_a.min_row}` = {rep(V(f'A{r_a.min_row}'))})"
                                              if r_a else "no está en un rango combinado") +
                      "; B" + f"{f} " + (f"dentro de `{r_b.coord}`." if r_b else "no está en un rango combinado.")]
            if f - 1 >= PRIMERA:
                partes.append(f"La fila {f - 1} " + (f"pertenece al rango `{r_arr.coord}` (principal `C{r_arr.min_row}` = {rep(V(f'C{r_arr.min_row}'))})"
                                                      if r_arr else "no está en un rango combinado de C") + ".")
            if arriba:
                iguales = all(V(f"{L}{f}") == V(f"{L}{arriba}") for L in "EFGH")
                partes.append(f"Código explícito previo más cercano: `C{arriba}` = {rep(V(f'C{arriba}'))}; "
                              f"E:H {'idénticos' if iguales else 'distintos'} a los de la fila {arriba}.")
            if abajo:
                partes.append(f"Siguiente código explícito: `C{abajo}` = {rep(V(f'C{abajo}'))}.")
            partes.append(f"Valores en D, I, J, K, L: " +
                          ("todos vacíos." if all(vacio(V(f'{L}{f}')) for L in 'DIJKL') else "hay contenido (ver tabla)."))
            w(" ".join(partes))
        w("")
        w("Nada en el libro permite determinar a qué servicio pertenecen; esa decisión no se toma aquí.")
    else:
        w("No hay filas sin código fuera de combinaciones.")
    w("")
    w("Para comparar, filas de continuación dentro de combinaciones de C (E:H con valor en todas ellas: " +
      ("sí" if all(all(not vacio(V(f"{L}{f}")) for L in "EFGH") for f in dentro) else "no") + "):")
    w("")
    w(", ".join(f"{f}" for f in dentro) + ".")
    w("")
    # filas totalmente vacías dentro de la zona
    vacias_zona = [f for f in range(PRIMERA, ULTIMA + 1) if all(vacio(V(f"{get_column_letter(i)}{f}")) for i in range(1, 13))]
    w("Filas completamente vacías en `A:L` dentro de " f"{PRIMERA}–{ULTIMA}: " +
      (rangos_compactos(vacias_zona) if vacias_zona else "ninguna") + ".")
    w("")

    # ============================================================ g
    w("## g. Listas de opciones `E112:H122`")
    w("")
    w(f"Fila de rótulos: " + ", ".join(f"`{L}{OPC_FILAS[0] - 1}` = {rep(V(f'{L}{OPC_FILAS[0] - 1}'))}" for L in "EFGH") + ".")
    w("")
    listas = {}
    for L in "EFGH":
        vals = [(f, V(f"{L}{f}")) for f in range(OPC_FILAS[0], OPC_FILAS[1] + 1) if V(f"{L}{f}") is not None]
        listas[L] = vals
        w(f"### Columna {L} — {OPCIONES[L]} ({len(vals)} valores)")
        w("")
        w(tabla(["Celda", "Valor exacto"], [[f"`{L}{f}`", rep(v)] for f, v in vals]))
        w("")
        sucios = [f"{L}{f}" for f, v in vals if isinstance(v, str) and v != v.strip()]
        w("Valores con espacios al inicio o al final: " + (", ".join(sucios) if sucios else "ninguno") + ".")
        w("")
    # nada fuera de E:H en el rango de opciones
    extra = [c.coordinate for row in ws.iter_rows(min_row=OPC_FILAS[0] - 1, max_row=OPC_FILAS[1])
             for c in row if c.value is not None and c.column not in (5, 6, 7, 8)]
    w("Celdas con valor en las filas 111–122 fuera de `E:H`: " + (", ".join(extra) if extra else "ninguna") + ".")
    w("")

    # ============================================================ h
    w("## h. Diferencias de etiquetas")
    w("")
    w("Se compara, columna por columna, la lista del Excel (sección g) con la del enunciado (sección 2). "
      "La columna E se compara con «Indicador S/N». Comparación exacta (distingue mayúsculas, espacios y acentos).")
    w("")
    hay_dif = False
    filas = []
    for L in "EFGH":
        ex = [v for _, v in listas[L]]
        en = enun_listas[L]
        solo_ex = [v for v in ex if v not in en]
        solo_en = [v for v in en if v not in ex]
        # pares que difieren solo en mayúsculas/espacios/acentos
        norm = lambda s: re.sub(r"\s+", " ", str(s)).strip().casefold()  # noqa: E731
        aprox = [(a, b) for a in solo_ex for b in solo_en if norm(a) == norm(b)]
        hay_dif |= bool(solo_ex or solo_en)
        orden = "mismo orden" if ex == en else ("distinto orden" if sorted(map(str, ex)) == sorted(map(str, en)) else "—")
        filas.append([f"{L} — {OPCIONES[L]}", len(ex), len(en),
                      ", ".join(rep(v) for v in solo_ex) or "ninguno",
                      ", ".join(rep(v) for v in solo_en) or "ninguno",
                      ", ".join(f"{rep(a)} ≈ {rep(b)}" for a, b in aprox) or "—", orden])
    w(tabla(["Columna", "N.º en Excel", "N.º en enunciado", "Solo en el Excel", "Solo en el enunciado",
             "Equivalentes salvo grafía", "Orden"], filas))
    w("")
    w("Resultado: " + ("**hay diferencias** (ver tabla)." if hay_dif else
                       "las cuatro listas del Excel y del enunciado contienen exactamente los mismos valores."))
    w("")
    # valores usados en los datos frente a las listas
    w(f"Valores efectivamente usados en las filas {PRIMERA}–{ULTIMA} (todas las celdas, incluidas las de "
      "continuación) frente a la lista de opciones del Excel:")
    w("")
    filas = []
    hay_fuera = False
    for L in "EFGH":
        usados = defaultdict(list)
        for f in range(PRIMERA, ULTIMA + 1):
            v = V(f"{L}{f}")
            if v is not None:
                usados[v].append(f)
        lista = [v for _, v in listas[L]]
        for v in sorted(usados, key=str):
            en_lista = v in lista
            hay_fuera |= not en_lista
            filas.append([f"{L} — {OPCIONES[L]}", rep(v), len(usados[v]), rangos_compactos(usados[v]),
                          "sí" if en_lista else "**no**"])
        n_vacias = sum(1 for f in range(PRIMERA, ULTIMA + 1) if vacio(V(f"{L}{f}")))
        filas.append([f"{L} — {OPCIONES[L]}", "*(vacía)*", n_vacias,
                      rangos_compactos([f for f in range(PRIMERA, ULTIMA + 1) if vacio(V(f"{L}{f}"))]) or "—", "—"])
    w(tabla(["Columna", "Valor", "Filas", "Cuáles", "En la lista del Excel"], filas))
    w("")
    w("Valores usados que no figuran en la lista: " + ("**sí, ver tabla**." if hay_fuera else "ninguno."))
    w("")
    # listas no usadas
    nunca = []
    for L in "EFGH":
        usados = {V(f"{L}{f}") for f in range(PRIMERA, ULTIMA + 1)}
        nunca += [f"{L}: {rep(v)}" for _, v in listas[L] if v not in usados]
    w("Opciones de la lista que no se usan en ninguna fila de datos: " + ("; ".join(nunca) if nunca else "ninguna") + ".")
    w("")

    # ============================================================ i
    w("## i. Mínimo y máximo (columnas K y L)")
    w("")
    for L in "KL":
        cnt = defaultdict(list)
        for f in range(PRIMERA, ULTIMA + 1):
            cnt[tipo(V(f"{L}{f}"))].append(f)
        w(f"### Columna {L} — {nombres[L]} (filas {PRIMERA}–{ULTIMA})")
        w("")
        w(tabla(["Tipo observado", "Cantidad", "Filas"],
                [[t, len(fs), rangos_compactos(fs)] for t, fs in sorted(cnt.items())]))
        w("")
        con_valor = [(f, V(f"{L}{f}")) for f in range(PRIMERA, ULTIMA + 1) if not vacio(V(f"{L}{f}"))]
        if con_valor:
            w("Celdas con valor: " + ", ".join(f"`{L}{f}` = {rep(v)}" for f, v in con_valor) + ".")
            w("")
        enrango = [r.coord for r in rangos if r.min_col <= (11 if L == "K" else 12) <= r.max_col]
        w(f"Rangos combinados que incluyen la columna {L}: " + (", ".join(enrango) if enrango else "ninguno") + ".")
        w("")
    pares = []
    for f in range(PRIMERA, ULTIMA + 1):
        k, l_ = V(f"K{f}"), V(f"L{f}")
        pares.append((f, k, l_))
    ambos = [(f, k, l_) for f, k, l_ in pares if not vacio(k) and not vacio(l_)]
    solo_k = [f for f, k, l_ in pares if not vacio(k) and vacio(l_)]
    solo_l = [f for f, k, l_ in pares if vacio(k) and not vacio(l_)]
    w("Combinaciones por fila:")
    w("")
    w(f"- Con K y L: {len(ambos)} fila(s) — " + (", ".join(
        f"{f} (K={rep(k)}, L={rep(l_)}" + (")" if isinstance(k, (int, float)) and isinstance(l_, (int, float)) and k <= l_
                                          else "; **mínimo > máximo o tipos no comparables**)") for f, k, l_ in ambos) or "—") + ".")
    w(f"- Solo K: {len(solo_k)} ({rangos_compactos(solo_k) or '—'}).")
    w(f"- Solo L: {len(solo_l)} ({rangos_compactos(solo_l) or '—'}).")
    w(f"- Ninguno: {len(pares) - len(ambos) - len(solo_k) - len(solo_l)} de {len(pares)} filas.")
    w("")
    w("Observación: los valores numéricos se leen como `float` (`1.0`, `100.0`); el formato de número "
      "mostrado en Excel no se determina aquí.")
    w("")

    # ============================================================ final
    w("## Texto sospechoso en los datos")
    w("")
    w("Dos revisiones: (1) todo el texto de la columna I (Descripción), única columna de texto libre, y de "
      "las celdas fuera de la tabla; (2) búsqueda por palabras clave en todas las celdas de texto. "
      "Ninguna se ha seguido; solo se citan.")
    w("")
    libres = [(f"I{f}", V(f"I{f}")) for f in range(PRIMERA, ULTIMA + 1) if V(f"I{f}") is not None]
    libres += [(c.coordinate, c.value) for c in fuera if isinstance(c.value, str)]
    w("Texto libre:")
    w("")
    w(tabla(["Celda", "Valor"], [[f"`{c}`", rep(v)] for c, v in libres]) if libres else "Ninguno.")
    w("")
    reg = re.compile("|".join(PATRONES_SOSPECHOSOS), re.I)
    claves = [(c.coordinate, c.value) for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and reg.search(c.value)]
    w("Coincidencias de palabras clave (" + ", ".join(f"`{p}`" for p in PATRONES_SOSPECHOSOS) + "):")
    w("")
    w(tabla(["Celda", "Valor"], [[f"`{c}`", rep(v)] for c, v in claves]) if claves else "Ninguna.")
    w("")
    w("Nota: `revel` se añadió a la lista tras leer `I5` (verbo en imperativo); la búsqueda por palabras "
      "clave no sustituye a la lectura de las celdas de texto libre listadas arriba.")
    w("")

    texto = "\n".join(out)
    SALIDA.write_text(texto, encoding="utf-8", newline="\n")
    print(f"Escrito {SALIDA.relative_to(RAIZ)} ({texto.count(chr(10))} líneas). N1={cn1} (esperado {ESPERADO_N1}), "
          f"N2={cn2} (esperado {ESPERADO_N2}), rangos combinados={len(rangos)}")


if __name__ == "__main__":
    main()
